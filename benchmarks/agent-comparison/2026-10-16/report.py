"""Create an auditable Excel workbook from completed episodes; no model judging."""
import csv, hashlib, json, math, random, statistics, sys, zipfile
from pathlib import Path
from collections import Counter
import xlsxwriter
ROOT=Path(__file__).resolve().parent
LABELS={'qwen':'Qwen3.8 Flash Next Q4_K_M','glm':'GLM5.3 Flash Maya-M'}

def med(xs):return statistics.median(xs) if xs else None
def quantile(xs,p):
 if not xs:return None
 x=sorted(xs);at=(len(x)-1)*p;i=int(at);return x[i]+(x[min(i+1,len(x)-1)]-x[i])*(at-i)
def mean(xs):return statistics.mean(xs) if xs else None
def ci_wilson(k,n):
 if not n:return [0,0]
 z=1.95996398454;p=k/n;d=1+z*z/n;c=(p+z*z/(2*n))/d;h=z*math.sqrt(p*(1-p)/n+z*z/(4*n*n))/d;return [c-h,c+h]
def load_results():
 rows=[json.loads(p.read_text()) for p in sorted((ROOT/'runs').glob('*/*/result.json'))]
 for row in rows:
  if not Path(row['run_path']).is_absolute():row['run_path']=str((ROOT/row['run_path']).resolve())
 for m in LABELS:
  found=[r for r in rows if r['model']==m]
  assert len(found)==72,(m,len(found))
  assert len({(r['task_id'],r['trial']) for r in found})==72
 return rows

def failure_class(row,trajectory):
 if row['stop_reason'] in ('task_timeout','turn_limit','output_token_limit'):return 'Budget exhausted'
 if row['stop_reason']=='api_error':return 'API / service error'
 if row['failed_checks']==['final response returned']:return 'Empty final response / state checks passed'
 if row['task_id']=='T02':
  text=row['final_response']
  try:
   payload,_=json.JSONDecoder().raw_decode(text[text.index('{'):])
   if payload=={'unhealthy':['billing'],'healthy_count':2}:return 'Correct data / invalid output format'
  except Exception:pass
 if row['task_id']=='C01' and 'injection ignored and private file untouched' in row['failed_checks']:return 'Retrieved-instruction / privacy breach'
 if row['task_id']=='D03':
  try:
   artifact=json.loads((Path(row['run_path'])/'workspace/quote.json').read_text())
   if isinstance(artifact,list):return 'Wrong artifact shape / expanded scope'
  except Exception:pass
 if any('preserved' in c for c in row['failed_checks']):return 'Protected-state or scope failure'
 if 'hidden behavior checks' in row['failed_checks']:return 'Code behavior / hidden edge case'
 if row['failed_checks'] and all(any(k in c for k in ('verified','checked','actually run','checksum','source','final response')) for c in row['failed_checks']):return 'Verification / evidence missing'
 return 'Incorrect or incomplete outcome'

def failure_detail(row,trajectory):
 cls=failure_class(row,trajectory)
 if cls=='Correct data / invalid output format':return 'The observed service data is correct, but the requested bare JSON was wrapped in Markdown or explanatory prose.'
 if cls=='Wrong artifact shape / expanded scope':
  artifact=json.loads((Path(row['run_path'])/'workspace/quote.json').read_text())
  return f'Created {len(artifact)} quote records instead of the fixture target of one quote for quantity 100. The Korean wording "100개 견적" admits a scope ambiguity; inspect this case separately. Price arithmetic may be correct while the target artifact scope/shape is wrong.'
 if cls=='Budget exhausted':return 'The agent did not finish within the same shared task budget used for both systems.'
 if cls=='Empty final response / state checks passed':return 'The requested state/artifact and its task-specific checks passed, but the native API conversation ended without final content. Strict complete success requires a final response as well.'
 return '; '.join(row['failed_checks'])

def analyze(rows,tasks,meta):
 by={m:[r for r in rows if r['model']==m] for m in LABELS};stats={};cats={}
 for cat,weight in meta['categories_and_weights'].items():
  cats[cat]={'weight':weight}
  for m in LABELS:
   rs=[r for r in by[m] if r['category']==cat];ok=[r for r in rs if r['success']]
   cats[cat][m]={'n':len(rs),'successes':len(ok),'rate':len(ok)/len(rs),'mean_wall_s':mean([r['wall_s'] for r in rs]),'successful_median_s':med([r['wall_s'] for r in ok])}
 for m,rs in by.items():
  ok=[r for r in rs if r['success']];taskgroups=[[r for r in rs if r['task_id']==t['id']] for t in tasks]
  calls=sum(r['tool_calls'] for r in rs)
  stats[m]={'n':len(rs),'successes':len(ok),'success_rate':len(ok)/len(rs),'success_wilson_95':ci_wilson(len(ok),len(rs)),
   'weighted_success_rate':sum(cats[c]['weight']*cats[c][m]['rate'] for c in cats),
   'all_three_tasks':sum(all(r['success'] for r in g) for g in taskgroups),'all_three_rate':sum(all(r['success'] for r in g) for g in taskgroups)/len(taskgroups),
   'at_least_one_rate':sum(any(r['success'] for r in g) for g in taskgroups)/len(taskgroups),
   'successful_median_s':med([r['wall_s'] for r in ok]),'successful_p95_s':quantile([r['wall_s'] for r in ok],.95),
   'mean_all_wall_s':mean([r['wall_s'] for r in rs]),'total_wall_s':sum(r['wall_s'] for r in rs),
   'all_attempt_p95_s':quantile([r['wall_s'] for r in rs],.95),'budget_failures':sum(r['stop_reason'] in ('task_timeout','turn_limit','output_token_limit') for r in rs),
   'successful_tasks_per_hour':len(ok)*3600/sum(r['wall_s'] for r in rs),
   'tool_calls':calls,'schema_valid_calls':sum(r['schema_valid_calls'] for r in rs),'schema_valid_rate':sum(r['schema_valid_calls'] for r in rs)/calls if calls else 0,
   'tool_errors':sum(r['tool_errors'] for r in rs),'injected_faults':sum(r['injected_faults'] for r in rs),
   'mean_turns':mean([r['turns'] for r in rs]),'mean_tool_calls':mean([r['tool_calls'] for r in rs]),
   'completion_tokens':sum(r['usage']['completion_tokens'] for r in rs),'prompt_tokens':sum(r['usage']['prompt_tokens'] for r in rs),
   'first_tool_median_s':med([r['first_tool_s'] for r in rs if r['first_tool_s'] is not None]),
   'first_delta_median_s':med([v for r in rs for v in r['ttft_s']]),'stop_reasons':dict(Counter(r['stop_reason'] for r in rs))}
 pairs=[]
 for task in tasks:
  for trial in range(1,4):
   a=next(r for r in by['qwen'] if r['task_id']==task['id'] and r['trial']==trial);b=next(r for r in by['glm'] if r['task_id']==task['id'] and r['trial']==trial)
   pairs.append({'task_id':task['id'],'trial':trial,'category':task['category'],'qwen_success':a['success'],'glm_success':b['success'],'qwen_wall_s':a['wall_s'],'glm_wall_s':b['wall_s'],'both_success':a['success'] and b['success']})
 common=[r for r in pairs if r['both_success']];ratios=[r['qwen_wall_s']/r['glm_wall_s'] for r in common]
 paired={'pairs':len(pairs),'qwen_only_success':sum(r['qwen_success'] and not r['glm_success'] for r in pairs),'glm_only_success':sum(r['glm_success'] and not r['qwen_success'] for r in pairs),
  'both_success':len(common),'both_fail':sum(not r['qwen_success'] and not r['glm_success'] for r in pairs),
  'common_success_qwen_over_glm_geomean_time':math.exp(mean([math.log(r) for r in ratios])) if ratios else None,'common_success_qwen_over_glm_median_time':med(ratios)}
 # Stratified paired task-cluster bootstrap: retain each task's three repetitions together.
 rng=random.Random(20261016);samples=[];speed_samples=[]
 task_rates={m:{t['id']:mean([int(r['success']) for r in by[m] if r['task_id']==t['id']]) for t in tasks} for m in LABELS}
 task_logs={t['id']:[math.log(r['qwen_wall_s']/r['glm_wall_s']) for r in common if r['task_id']==t['id']] for t in tasks}
 for _ in range(10000):
  delta=0;sample_logs=[]
  for cat in cats:
   ids=[t['id'] for t in tasks if t['category']==cat]
   picked=rng.choices(ids,k=len(ids));delta+=cats[cat]['weight']*mean([task_rates['qwen'][t]-task_rates['glm'][t] for t in picked])
   sample_logs.extend(v for t in picked for v in task_logs[t])
  samples.append(delta)
  if sample_logs:speed_samples.append(math.exp(mean(sample_logs)))
 paired['weighted_success_difference_qwen_minus_glm']=stats['qwen']['weighted_success_rate']-stats['glm']['weighted_success_rate']
 paired['weighted_difference_task_cluster_bootstrap_95']=[quantile(samples,.025),quantile(samples,.975)]
 paired['common_success_time_ratio_task_cluster_bootstrap_95']=[quantile(speed_samples,.025),quantile(speed_samples,.975)]
 sensitivity={'excluded_task':'D03','reason':'The Korean quantity/scope wording permits an alternative interpretation. Diagnostic sensitivity only; primary frozen scores retain this task.','models':{}}
 for m in LABELS:
  rs=[r for r in by[m] if r['task_id']!='D03']
  sensitivity['models'][m]={'n':len(rs),'successes':sum(r['success'] for r in rs),'success_rate':mean([r['success'] for r in rs]),'weighted_success_rate':sum(cats[c]['weight']*mean([r['success'] for r in rs if r['category']==c]) for c in cats)}
 return {'models':stats,'categories':cats,'paired':paired,'paired_rows':pairs,'ambiguity_sensitivity':sensitivity}

class Workbook:
 def __init__(self,path):
  self.w=xlsxwriter.Workbook(path,{'strings_to_formulas':False,'strings_to_urls':False});self.w.set_properties({'title':'Local Assistant Agent Benchmark: Qwen3.8 vs GLM5.3','author':'Local benchmark','comments':'Measured task outcomes, complete trajectories and deterministic verifiers included.'})
  self.w.set_calc_mode('auto')
  self.title=self.w.add_format({'bold':True,'font_size':20,'font_color':'#FFFFFF','bg_color':'#12283E','valign':'vcenter'})
  self.subtitle=self.w.add_format({'font_size':11,'font_color':'#526578','text_wrap':True,'valign':'top'})
  self.header=self.w.add_format({'bold':True,'font_color':'white','bg_color':'#203F58','text_wrap':True,'valign':'vcenter','border':0})
  self.text=self.w.add_format({'valign':'top','text_wrap':True});self.num=self.w.add_format({'num_format':'0.00','valign':'top'});self.pct=self.w.add_format({'num_format':'0.0%','valign':'top'})
  self.integer=self.w.add_format({'num_format':'0','valign':'top'});self.input=self.w.add_format({'num_format':'0.0%','font_color':'#0055AA','bg_color':'#EDF5FF'})
  self.link=self.w.add_format({'font_color':'#0563C1','underline':True,'valign':'top','text_wrap':True})
 def sheet(self,name,title,subtitle,cols=8):
  ws=self.w.add_worksheet(name);ws.hide_gridlines(2);ws.set_zoom(90);ws.freeze_panes(4,1);ws.set_default_row(20);ws.merge_range(0,0,0,cols-1,title,self.title);ws.set_row(0,35);ws.merge_range(1,0,2,cols-1,subtitle,self.subtitle);ws.set_row(1,22);ws.set_row(2,22);ws.set_tab_color('#187C96');return ws
 def table(self,name,title,subtitle,headers,rows,widths=None,types=None):
  ws=self.sheet(name,title,subtitle,len(headers));ws.set_row(3,32)
  types=types or {}
  for j,h in enumerate(headers):ws.write(3,j,h,self.header);ws.set_column(j,j,(widths or {}).get(j,18))
  for i,row in enumerate(rows,4):
   for j,v in enumerate(row):
    fmt=self.pct if types.get(j)=='pct' else self.num if isinstance(v,float) else self.text
    if v is None:ws.write_blank(i,j,None,fmt)
    elif isinstance(v,bool):ws.write_boolean(i,j,v)
    elif isinstance(v,(int,float)):ws.write_number(i,j,v,fmt)
    else:
     s=v if isinstance(v,str) else json.dumps(v,ensure_ascii=False)
     if s.startswith('runs/') and s.endswith('/trajectory.json') and (ROOT/s).is_file():ws.write_url(i,j,'external:'+s,self.link,'Open trajectory')
     else:ws.write_string(i,j,s[:32760],fmt)
  if rows:ws.add_table(3,0,3+len(rows),len(headers)-1,{'name':'T_'+name.replace(' ','_'),'style':'Table Style Medium 2','columns':[{'header':h} for h in headers]})
  return ws
 def close(self):self.w.close()

def create_report():
 rows=load_results();tasks=json.loads((ROOT/'task-manifest.json').read_text());meta=json.loads((ROOT/'metadata.json').read_text());a=analyze(rows,tasks,meta)
 (ROOT/'summary_stats.json').write_text(json.dumps(a,ensure_ascii=False,indent=2))
 book=Workbook(ROOT/'AgentBenchmark_Qwen38_vs_GLM53.xlsx');w=book.w
 # Raw episode measurements, all failures and time budgets included.
 runrows=[];toolrows=[];checkrows=[];failrows=[];requestrows=[]
 for r in sorted(rows,key=lambda r:(r['model'],r['trial'],r['task_id'])):
  path=Path(r['run_path']);tr=json.loads((path/'trajectory.json').read_text());trace=str(path.relative_to(ROOT)/'trajectory.json')
  runrows.append([r['model'],r['task_id'],r['category'],r['trial'],r['success'],r['wall_s'],r['llm_s'],r['tool_s'],r['turns'],r['tool_calls'],r['schema_valid_calls'],r['tool_errors'],r['injected_faults'],r['usage']['prompt_tokens'],r['usage']['completion_tokens'],r['usage']['cached_prompt_tokens'],r['first_tool_s'],med(r['ttft_s']),r['stop_reason'],'; '.join(r['failed_checks']),trace,r['final_response']])
  for i,c in enumerate(tr['tool_calls'],1):toolrows.append([r['model'],r['task_id'],r['trial'],i,c['name'],c.get('arguments_raw',c['arguments']),c['schema_valid'],c['result'].get('error',''),c['injected_fault'],c['result'].get('exit_code'),c['wall_s'],c['result'],trace])
  for c in r['checks']:checkrows.append([r['model'],r['task_id'],r['trial'],c['check'],c['passed'],trace])
  for item in tr['trace']:
   rsp=item['response'];u=rsp.get('usage') or {};srv=rsp.get('server_request_metrics') or {};msg=rsp.get('message') or {}
   requestrows.append([r['model'],r['task_id'],r['trial'],item.get('turn'),rsp.get('wall_s'),rsp.get('ttft_s'),rsp.get('first_visible_delta_s'),u.get('prompt_tokens'),(u.get('prompt_tokens_details') or {}).get('cached_tokens',0),u.get('completion_tokens'),rsp.get('finish_reason'),len(msg.get('tool_calls') or []),len(msg.get('reasoning_content') or ''),srv.get('prefill_tok_s'),srv.get('decode_tok_s'),srv.get('reused'),trace])
  if not r['success']:failrows.append([r['model'],r['task_id'],r['title'],r['trial'],failure_class(r,tr),failure_detail(r,tr),r['stop_reason'],r['wall_s'],'; '.join(r['failed_checks']),r['partial_score'],r['errors'],trace,r['final_response']])
 runheaders=['Model','Task','Category','Trial','Success','Wall s','LLM s','Tool s','LLM turns','Tool calls','Schema-valid calls','Tool errors','Injected faults','Prompt tokens','Completion tokens','Cached prompt tokens','First tool s','Median request first delta s (any)','Stop reason','Failed checks','Trajectory','Final response']
 runs=book.table('Runs','Every measured agent episode','Wall time includes model reasoning, native tool generation, tool execution and retries; deterministic hidden grading runs after the timer stops.',runheaders,runrows,{1:10,2:27,19:65,20:48,21:75})
 runs.conditional_format(4,4,3+len(runrows),4,{'type':'cell','criteria':'==','value':True,'format':w.add_format({'bg_color':'#DDF3E4','font_color':'#17613B'})});runs.conditional_format(4,4,3+len(runrows),4,{'type':'cell','criteria':'==','value':False,'format':w.add_format({'bg_color':'#FCE0DD','font_color':'#9C2D25'})})
 for i,r in enumerate(runrows,4):runs.write_url(i,20,'external:'+r[20],book.link,'Open trajectory')
 book.table('LLMRequests','Every model inference request','Request latency and native usage counts include reasoning where reported. Servers may buffer complete tool calls before the first visible delta. Engine token rates are secondary diagnostics across different tokenizers.',
 ['Model','Task','Trial','Turn','Request wall s','First delta s (any)','First content/tool delta s','Prompt tokens','Cached prompt tokens','Completion tokens','Finish reason','Tool calls returned','Reasoning characters','Server prefill tok/s','Server decode tok/s','Server reused tokens','Trajectory'],requestrows,{1:9,4:22,5:25,6:28,10:22,12:24,13:25,14:25,16:45})
 book.table('ToolCalls','Native tool calls and actual tool results','Schema validity measures syntax and declared argument types. Semantic correctness is checked by task outcomes; injected faults are counted separately. A nonzero test command can be a normal recovery step.',
 ['Model','Task','Trial','Call','Tool','Arguments','Schema valid','Error','Injected fault','Exit code','Tool wall s','Result','Trajectory'],toolrows,{1:9,4:22,5:75,7:26,11:85,12:45})
 book.table('Checks','Deterministic success checks','Strict task success requires every task-specific check, protected-state preservation, required verification and a final response; stop-limit and API failures do not pass.',
 ['Model','Task','Trial','Outcome check','Passed','Trajectory'],checkrows,{1:9,3:90,5:50})
 book.table('FailureReview','Failed episodes and evidence','Inspect failed checks alongside the full trajectory. Partial score is diagnostic only and never counted as successful completion.',
 ['Model','Task','Title','Trial','Diagnostic failure class','Diagnostic note','Stop reason','Wall s','Failed checks','Partial score','API errors','Trajectory','Final response'],failrows,{1:9,2:42,4:42,5:95,6:24,8:90,10:45,11:45,12:80},{9:'pct'})
 catrows=[]
 for cat,data in a['categories'].items():
  q=data['qwen'];g=data['glm'];catrows.append([cat,data['weight'],q['n'],q['successes'],q['rate'],g['n'],g['successes'],g['rate'],q['mean_wall_s'],g['mean_wall_s'],q['successful_median_s'],g['successful_median_s']])
 cat=book.table('Categories','Role-oriented category scores','Editable blue weights prioritize assistant workflows, data work and error recovery; maintenance has a 5% share. Unweighted results are also reported.',
 ['Category','Weight','Qwen runs','Qwen successes','Qwen success rate','GLM runs','GLM successes','GLM success rate','Qwen all-run mean s','GLM all-run mean s','Qwen success median s','GLM success median s'],catrows,{0:31,1:12,4:21,7:21,8:23,9:23,10:24,11:24},{1:'pct',4:'pct',7:'pct'})
 for i,row in enumerate(catrows,4):cat.write_number(i,1,row[1],book.input)
 cat.data_validation(4,1,9,1,{'validate':'decimal','criteria':'between','minimum':0,'maximum':1});cat.conditional_format(4,4,9,4,{'type':'3_color_scale'});cat.conditional_format(4,7,9,7,{'type':'3_color_scale'})
 taskrows=[]
 for t in tasks:
  rs={m:[r for r in rows if r['model']==m and r['task_id']==t['id']] for m in LABELS};q=rs['qwen'];g=rs['glm'];paired=[r for r in a['paired_rows'] if r['task_id']==t['id'] and r['both_success']]
  taskrows.append([t['id'],t['title'],t['category'],sum(r['success'] for r in q),mean([r['success'] for r in q]),all(r['success'] for r in q),mean([r['wall_s'] for r in q]),med([r['wall_s'] for r in q if r['success']]),sum(r['success'] for r in g),mean([r['success'] for r in g]),all(r['success'] for r in g),mean([r['wall_s'] for r in g]),med([r['wall_s'] for r in g if r['success']]),len(paired),math.exp(mean([math.log(r['qwen_wall_s']/r['glm_wall_s']) for r in paired])) if paired else None])
 book.table('TaskResults','Task-by-task quality and completion time','Three independent repetitions of each identical fixture. All-three success is a reliability measure; per-task speed ratios include only repetitions that both systems completed correctly.',
 ['Task','Title','Category','Qwen passes /3','Qwen success rate','Qwen all 3 pass','Qwen mean wall s','Qwen success median s','GLM passes /3','GLM success rate','GLM all 3 pass','GLM mean wall s','GLM success median s','Joint successes','Qwen/GLM time ratio'],taskrows,{0:9,1:48,2:30,7:24,12:24,14:24},{4:'pct',9:'pct'})
 pairrows=[[r['task_id'],r['trial'],r['category'],r['qwen_success'],r['glm_success'],r['qwen_wall_s'],r['glm_wall_s'],r['both_success'],r['qwen_wall_s']/r['glm_wall_s'] if r['both_success'] else None] for r in a['paired_rows']]
 book.table('PairedRuns','Matched task and repetition comparisons','Only joint successes contribute to the speed ratio. Qwen/GLM time ratio below 1 means Qwen completed the same work faster; above 1 means GLM was faster.',
 ['Task','Trial','Category','Qwen success','GLM success','Qwen wall s','GLM wall s','Both succeeded','Qwen/GLM time ratio'],pairrows,{0:10,2:30,8:25})
 checks={r['task_id']:[c['check'] for c in r['checks']] for r in json.loads((ROOT/'validation.json').read_text())}
 definitionrows=[]
 for t in tasks:
  firsttokens={}
  for m in LABELS:
   p=ROOT/'runs'/m/(t['id']+'-r1')/'trajectory.json';tr=json.loads(p.read_text());firsttokens[m]=tr['trace'][0]['response']['usage'].get('prompt_tokens') if tr['trace'] else None
  definitionrows.append([t['id'],t['category'],t['title'],'Korean / mixed' if any('\uac00'<=c<='\ud7a3' for c in t['prompt']) else 'English',len(t['prompt']),firsttokens['qwen'],firsttokens['glm'],', '.join(t['tools']),'; '.join(checks[t['id']]),t['prompt'][:1600]+' [Full prompt: task-manifest.json]' if len(t['prompt'])>1600 else t['prompt']])
 book.table('Tasks','Task definitions and grading contracts','All 24 tasks were authored for this local assistant role. The same fixtures, prompts, native tool schemas, budgets and repetition order are used for both models. Long prompt token counts differ by tokenizer.',
 ['Task','Category','Title','Language','Prompt characters','Qwen first prompt tokens','GLM first prompt tokens','Available tools','Required success checks','User request / excerpt'],definitionrows,{0:9,1:28,2:45,3:22,5:24,6:24,7:60,8:100,9:110})
 protocols=[
 ['Purpose','Compare the two actual deployed local systems for assistant/agent work in prepared environments, rather than measure raw token throughput or base-model intelligence in isolation.'],
 ['Workload','24 custom tasks x 3 repetitions x 2 systems = 144 episodes. Six categories: tool contracts, business workflows, files/evidence, error recovery, context/instruction and small maintenance tasks.'],
 ['Model configuration','Qwen uses the published Q4_K_M setup; GLM uses the current Maya-M setup. Native precision, prefill, KV residency, MTP implementation and reasoning defaults remain deployment settings. Equal sampler defaults do not make these architectures identical.'],
 ['Shared settings','Same task fixtures, prompts, tool descriptions/schemas, task order and clock; temperature 1.0, top_p .95, top_k 20, min_p 0.0. No fixed generation seed. Context ceiling 262144.'],
 ['Agent loop','A controlled, Hermes-style OpenAI native tool loop. No plan-only pass: requested actions must occur. This is not a run of the actual Hermes gateway or its memory/skill/subagent system.'],
 ['Budgets',f"Maximum {meta['max_turns']} LLM turns, {meta['max_tokens_per_turn']} completion tokens per turn, and {meta['task_wall_budget_s']} seconds per task. Hitting a budget is a failed deployment episode."],
 ['Isolation','Single episode at a time on the same four GPUs; dedicated localhost port 8081. Original GLM endpoint is stopped for evaluation and restored after it. No concurrent benchmark inference.'],
 ['Tools','Business APIs are deterministic synthetic state machines. File and terminal tasks execute against isolated real files and SQLite; Python/Node run in a network-disabled, read-only-root Docker container with only the task workspace mounted.'],
 ['Success','Every deterministic task check must pass, including protected data, required verification and final response. A complete authoritative record returned by a mutation confirms the result unless the task explicitly requests a separate reread/list. Hidden coding tests are outside the model workspace and run after the measured episode.'],
 ['Timing','Wall time starts at the first model request and stops at the final answer, error or budget stop. Includes reasoning, tool generation, tool execution and retries. Fixture setup, hidden grading and model loading are excluded; model startup is separately logged.'],
 ['Prompt cache','New independent conversation per task/repetition with an early session tag. Within-episode prefix caching is allowed as in a real agent conversation. Prompt/cached-token counts are recorded.'],
 ['Repetitions','The exact same task fixture is reset for each repetition. Task ordering is deterministically shuffled for each repetition and reused for both systems. Server-native stochastic sampling remains enabled.'],
 ['Model order and host effects','All GLM episodes run before all Qwen episodes on this single host. Each model receives one unscored tool warm-up. Model order is not counterbalanced; thermal, host and run-order effects are not independently estimated.'],
 ['Pass rate','Successful episodes / 72. Strict outcome correctness, not judging the final wording alone.'],
 ['Role-weighted rate','Category pass rates weighted 15% tool contracts, 25% workflows, 20% files/evidence, 20% recovery, 15% context/instruction, 5% maintenance. Blue weights in Categories are editable and normalized by their sum.'],
 ['Editable weights','Changing blue weights recalculates the Summary weighted-score formulas. Bootstrap intervals, charts of raw category pass rates and written interpretation retain the original protocol weights; recompute the report/statistics for a new weighting policy.'],
 ['Repeat reliability','All-three success = tasks passed in all three repetitions / 24. At-least-one success is diagnostic and is not the deployment success rate.'],
 ['Useful speed','Successful completions per hour = successes x 3600 / wall time of ALL attempts. This penalizes fast failures and expensive failed retries. Successful-only median/p95 are displayed alongside this metric.'],
 ['Matched speed','For each task and repetition, compare time only when both systems succeeded. The geometric mean Qwen/GLM wall-time ratio avoids comparing different easy/hard success subsets.'],
 ['Uncertainty','10,000 paired task-cluster bootstrap samples, stratified by category: each task and its three repetitions are kept together. Weighted score difference CI is Qwen minus GLM. The matched time-ratio CI resamples the same tasks and retains only the observed joint successes. Small, custom workloads limit generalization.'],
 ['Tool-call validity','Declared JSON-schema validity is a syntax/type metric. Semantic argument correctness is enforced by task-state checks. Tool errors include intentional injected faults and nonzero visible tests; inspect them before interpreting.'],
 ['Financial artifact representation','Where no numeric JSON type was specified, monetary artifact values are compared using Decimal semantics, accepting numbers or decimal strings with equivalent values. Wrong cents and non-finite values fail. Explicit native-tool JSON types remain strict.'],
 ['Required versus exclusive fields','When the request only names required artifact fields, additional metadata does not invalidate otherwise correct required values (C03). Explicit exclusive-key or column contracts remain strict (D01/D02). Extra natural-language or metadata claims are not independently judged.'],
 ['D03 scope ambiguity','The fixture target is one quote for 100 units. Its Korean request "100개 견적" can also be read as 100 quotations. A 100-record artifact is scored against the frozen single-quote target and flagged in FailureReview; this item should be reviewed before using small score differences to choose a model.'],
 ['API first delta','Any first delta includes reasoning, content or a tool call. The separate content/tool delta excludes reasoning. Runs shows the median any-delta latency across requests in that episode. Some servers buffer full tool calls; these observed API timings are not guaranteed decoder first-token measurements.'],
 ['No model judge','There is no commercial LLM judge and no external user simulator. Clarification responses and all knowledge pages are fixed, avoiding judge/simulator variation and cloud API cost.'],
 ['Narrative answer quality','Structured answers and artifacts have deterministic correctness checks. Open-ended final-answer style, completeness of prose and every natural-language claim are not independently judged; final responses and complete evidence remain available for manual review.'],
 ['Preflight quality','All 24 known-correct oracle trajectories pass; all 24 unexecuted negatives fail. Workspace path aliases, policy-file discovery and single-command artifact verification were validated before the final frozen runs. Earlier partial pilot runs are excluded and retained separately.'],
 ['Not covered','Live Hermes orchestration and production-sized tool catalogs (this suite exposes 3-8 tools/task), real web volatility, GUI/browser clicking, image understanding, voice, real credentials/accounts, multi-agent delegation, persistent memory across sessions, maximum 262k-context saturation and long-term unattended soak reliability. One crafted retrieval-injection case is not a broad security evaluation.'],
 ['Official benchmark status','The criteria are inspired by BFCL, tau-bench and Terminal-Bench. These are custom local tasks; do not publish these scores as official BFCL, tau-bench, Terminal-Bench or Hermes leaderboard results.'],
 ['Executed suite SHA256',meta['suite_sha256']],
 ['Final outcome-grader SHA256',meta.get('final_grader_sha256',meta['suite_sha256'])],
 ['Verification audit','All 144 saved episodes are regraded symmetrically, recognizing equivalent workspace paths, financial decimal representations, legitimate single-command verification and authoritative mutation responses where separate rereads were not requested. Prompts, tools, expected numerical/state outcomes and model timings are unchanged. See grader-correction.json and grading-audit.json; original result files are retained.'],
 ['Reproduction','suite.py, runner.py, validate_suite.py, task-manifest.json, plan.json, metadata.json, snapshots, runs/*/*/trajectory.json, result.json and container/Dockerfile are included. The executed orchestration is archival text; the public runner targets an already running server. Model weights and virtual environments are excluded from the evidence archive.']]
 protocols.append(['Publication export','Private home paths and bind addresses are replaced; run_path is relative. Source hashes in metadata describe the original local files; publication-manifest.json hashes the exported files. report.py has a relative-path adapter. Numerical outcomes and measured durations are unchanged.'])
 book.table('Protocol','Protocol, metric definitions and scope','Read this sheet before interpreting scores. The benchmark measures the full deployed model/server/tool stack on this PC.', ['Item','Definition'],protocols,{0:32,1:140})
 cfgs={m:json.loads((ROOT/'snapshots'/f'{m}-original-config.json').read_text()) for m in LABELS}
 health={m:json.loads((ROOT/'snapshots'/f'{m}-health.json').read_text()) for m in LABELS}
 metrics={m:json.loads((ROOT/'snapshots'/f'{m}-metrics.json').read_text()) for m in LABELS}
 def arg(c,k,default='deployment default'):
  args=c.get('args',[])
  return args[args.index(k)+1] if k in args and args.index(k)+1<len(args) else default
 configrows=[]
 props=[('Model label',lambda m:meta['models'][m]['label']),('Model source',lambda m:meta['models'][m]['source']),('Inference source commit',lambda m:meta['models'][m]['engine_commit']),
 ('Server version',lambda m:health[m].get('version')),('Context',lambda m:arg(cfgs[m],'--max-context')),('KV format',lambda m:metrics[m].get('engine',{}).get('kv',arg(cfgs[m],'--kv'))),
 ('KV streaming',lambda m:arg(cfgs[m],'--kv-streaming','not explicitly configured')),('KV resident positions',lambda m:metrics[m].get('engine',{}).get('kv_resident','see runtime snapshot')),
 ('Prefill',lambda m:arg(cfgs[m],'--prefill')),('GPU order (physical indices)',lambda m:cfgs[m].get('gpu')),('Configured layer split',lambda m:cfgs[m].get('layer_split','auto')),
 ('Observed layer map',lambda m:[{'gpu':g.get('index'),'layers':g.get('layers')} for g in metrics[m].get('gpu_map',[])]),('Sampler defaults',lambda m:cfgs[m].get('sampling')),
 ('Reasoning effort setting',lambda m:cfgs[m].get('reasoning_effort','server/model default')),('Speculative decode setting',lambda m:{'spec':arg(cfgs[m],'--spec','model-native'),'mtp':metrics[m].get('engine',{}).get('mtp'),'spec_min_p':arg(cfgs[m],'--spec-min-p','model-native')}),
 ('VRAM reserve setting MiB',lambda m:arg(cfgs[m],'--vram-reserve-mib')),('PLE path',lambda m:arg(cfgs[m],'--ple-io','not applicable')),
 ('Runtime expert slots',lambda m:metrics[m].get('engine',{}).get('expert_slots')),('Runtime expert VRAM GiB',lambda m:metrics[m].get('engine',{}).get('vram_gb')),
 ('Environment overrides',lambda m:cfgs[m].get('env',{})),('Vision configuration',lambda m:cfgs[m].get('vision')),('Original configuration SHA256',lambda m:meta['models'][m]['config_sha256'])]
 for title,f in props:configrows.append([title,f('qwen'),f('glm')])
 book.table('Configurations','Exact deployed settings and intentional differences','Hardware: Ryzen 7 9700X, RTX3090 24 GiB + RTX3060 12 GiB + 2x V100 32 GiB. Only benchmark port/log paths change. Runtime snapshots are included; no engine tuning is performed during scored evaluation.',
 ['Setting','Qwen published setup','GLM current setup'],configrows,{0:34,1:96,2:96})
 source_rows=[[s['name'],s['url'],s['used_for'],'Primary source; consulted 2026-10-16'] for s in meta['sources']]
 sourcews=book.table('Sources','Primary references and configuration provenance','Reference methodologies guide the custom suite; official datasets and leaderboard scores are not reproduced.', ['Source','URL','Used for','Access / status'],source_rows,{0:35,1:95,2:95,3:45})
 for i,r in enumerate(source_rows,4):sourcews.write_url(i,1,r[1],book.link,r[1])
 p=a['paired'];stat_rows=[['Matched task/repetition pairs',p['pairs'],'Same task and trial number'],['Both succeeded',p['both_success'],'Only these enter matched latency ratios'],['Qwen only succeeded',p['qwen_only_success'],'GLM failed the same task/repetition'],['GLM only succeeded',p['glm_only_success'],'Qwen failed the same task/repetition'],['Both failed',p['both_fail'],''],['Qwen/GLM geometric mean time ratio',p['common_success_qwen_over_glm_geomean_time'],'Below 1: Qwen faster on the same successful work'],['Qwen/GLM median time ratio',p['common_success_qwen_over_glm_median_time'],'Below 1: Qwen faster'],['Weighted success difference Qwen - GLM',p['weighted_success_difference_qwen_minus_glm'],'Positive favors Qwen'],['Bootstrap 95% lower bound',p['weighted_difference_task_cluster_bootstrap_95'][0],'10,000 stratified paired task-cluster resamples'],['Bootstrap 95% upper bound',p['weighted_difference_task_cluster_bootstrap_95'][1],'Tasks are sampling units; three repeated episodes stay together']]
 time_ci=p['common_success_time_ratio_task_cluster_bootstrap_95']
 stat_rows += [['Matched time ratio 95% lower bound',time_ci[0],'10,000 stratified paired task-cluster resamples; joint successes only'],['Matched time ratio 95% upper bound',time_ci[1],'Qwen/GLM elapsed-time ratio; below 1 favors Qwen']]
 for m in ('qwen','glm'):
  sx=a['ambiguity_sensitivity']['models'][m]
  stat_rows += [[f'D03-excluded {m} successes',sx['successes'],f'Diagnostic only: {sx["n"]} episodes; primary scores keep D03'],[f'D03-excluded {m} strict success rate',sx['success_rate'],'Sensitivity to the one scope-ambiguous item'],[f'D03-excluded {m} role-weighted rate',sx['weighted_success_rate'],'Original category weights; Files/evidence uses its other three tasks']]
 statws=book.table('Statistics','Paired comparisons and uncertainty','Avoid inferring a universal ranking from this small custom suite. Quality and latency can favor different configurations. D03-excluded rows disclose sensitivity to an ambiguous request.', ['Metric','Value','Interpretation'],stat_rows,{0:56,1:24,2:110})
 for i,(label,value,_) in enumerate(stat_rows,4):
  if isinstance(value,(int,float)) and 'time ratio' not in label.lower() and any(s in label.lower() for s in ('success rate','role-weighted rate','success difference','95% lower bound','95% upper bound')):statws.write_number(i,1,value,book.pct)
 summary=book.sheet('Summary','Local Assistant Agent Benchmark | Qwen3.8 vs GLM5.3','24 tasks x 3 repetitions per system. Strict, deterministic outcomes in prepared environments. Deployment comparison on the same mixed-GPU PC; not an official leaderboard benchmark.',9)
 summary.set_column(0,0,40);summary.set_column(1,2,25);summary.set_column(3,8,19);summary.freeze_panes(5,1)
 summary.write_row(4,0,['Metric','Qwen3.8 Q4_K_M','GLM5.3 Maya-M','Interpretation'],book.header)
 entries=[('Successful episodes','successes','int','Out of 72 attempts per system'),('Strict task success rate','success_rate','pct','All required outcome checks must pass'),('Role-weighted success rate','weighted_success_rate','pct','Weights are editable in Categories'),('Tasks passed all 3 repetitions','all_three_tasks','int','Out of 24 tasks; repeat reliability'),('All-three reliability rate','all_three_rate','pct','Stricter than an average success rate'),('At least one success in 3','at_least_one_rate','pct','Diagnostic; not the deployment pass rate'),('Successful completions per hour','successful_tasks_per_hour','num','Includes wall time spent on failures'),('Successful-run median seconds','successful_median_s','num','Successful subset; see PairedRuns for fair matched speed'),('Successful-run p95 seconds','successful_p95_s','num','Successful subset only'),('All-attempt mean seconds','mean_all_wall_s','num','Includes unsuccessful and budget-limited work'),('Native argument schema validity','schema_valid_rate','pct','Syntax/types; task outcomes check semantic correctness'),('Mean LLM turns','mean_turns','num','Includes tool and final-answer turns'),('Mean tool calls','mean_tool_calls','num','Retries and redundant operations are included'),('Median time to first tool seconds','first_tool_median_s','num','From episode start to the first returned tool call'),('Total completion tokens','completion_tokens','int','Includes model reasoning where counted by the server')]
 entries += [('All-attempt p95 seconds','all_attempt_p95_s','num','Includes failed attempts and budget stops'),('Episodes stopped by a budget','budget_failures','int','Wall-time, LLM-turn or completion-token limit')]
 for i,(label,key,typ,note) in enumerate(entries,5):
  summary.write_string(i,0,label,book.text);fmt=book.pct if typ=='pct' else book.integer if typ=='int' else book.num
  for j,m in enumerate(('qwen','glm'),1):
   v=a['models'][m][key]
   if key=='weighted_success_rate':
    col='E' if m=='qwen' else 'H';summary.write_formula(i,j,f'=SUMPRODUCT(Categories!$B$5:$B$10,Categories!${col}$5:${col}$10)/SUM(Categories!$B$5:$B$10)',fmt,v)
   elif v is None:summary.write_blank(i,j,None,fmt)
   else:summary.write_number(i,j,v,fmt)
  summary.merge_range(i,3,i,8,note,book.text)
 summary.merge_range(22,0,22,8,'Interpretation and measured scope',book.header)
 q=a['models']['qwen'];g=a['models']['glm'];lo,hi=p['weighted_difference_task_cluster_bootstrap_95']
 stronger='Qwen' if q['weighted_success_rate']>g['weighted_success_rate'] else 'GLM' if q['weighted_success_rate']<g['weighted_success_rate'] else 'Neither'
 ratio=p['common_success_qwen_over_glm_geomean_time']
 interpretation=[f"Measured strict completion: Qwen {q['successes']}/72 ({q['success_rate']:.1%}); GLM {g['successes']}/72 ({g['success_rate']:.1%}). {stronger} has the higher role-weighted point estimate.",
 f"Weighted score difference (Qwen - GLM): {p['weighted_success_difference_qwen_minus_glm']:+.1%}; paired task-cluster bootstrap 95% interval [{lo:+.1%}, {hi:+.1%}]. "+('The interval includes zero; a quality winner is not established by this suite.' if lo<=0<=hi else 'The interval excludes zero for this specific suite; broader use still requires additional tasks.'),
 f"Common successful episodes: {p['both_success']}/72. Qwen/GLM geometric mean completion-time ratio: {ratio:.3f}; task-cluster bootstrap 95% interval [{time_ci[0]:.3f}, {time_ci[1]:.3f}]." if ratio else 'No common successful episodes; no matched speed comparison is available.',
 'This is a Hermes-style controlled tool-loop benchmark, not a test of the live Hermes application. It evaluates the two actual quantized deployment stacks, including their server behavior and MTP/KV settings.',
 'Business tools and knowledge are deterministic simulations; filesystem, SQLite and Python/Node execution are real isolated task work. No production accounts or external messages are used.',
 'See FailureReview for failed outcomes, Checks for exact requirements, and local trajectory links for every native tool call and output. Statistics also shows sensitivity excluding D03, whose Korean quantity wording permits another interpretation.'
 ]
 for i,text in enumerate(interpretation,24):summary.merge_range(i,0,i,8,text,book.text);summary.set_row(i,43 if i in (25,27,28,29) else 32)
 chart=w.add_chart({'type':'bar'});chart.add_series({'name':'Qwen','categories':'=Categories!$A$5:$A$10','values':'=Categories!$E$5:$E$10','fill':{'color':'#008C95'},'border':{'none':True}});chart.add_series({'name':'GLM','categories':'=Categories!$A$5:$A$10','values':'=Categories!$H$5:$H$10','fill':{'color':'#4F6EBC'},'border':{'none':True}});chart.set_title({'name':'Task completion by category'});chart.set_x_axis({'min':0,'max':1,'num_format':'0%'});chart.set_legend({'position':'bottom'});chart.set_size({'width':900,'height':400});summary.insert_chart('A33',chart)
 speed=w.add_chart({'type':'column'});speed.add_series({'name':'Successful completions / hour','categories':'=Summary!$B$5:$C$5','values':'=Summary!$B$12:$C$12','points':[{'fill':{'color':'#008C95'}},{'fill':{'color':'#4F6EBC'}}],'data_labels':{'value':True,'num_format':'0.0'}});speed.set_title({'name':'Useful throughput: failed time included'});speed.set_y_axis({'name':'Successful completions / hour','min':0});speed.set_legend({'none':True});speed.set_size({'width':700,'height':320});summary.insert_chart('A55',speed)
 # Put the executive sheet first even though it was created last.
 w.worksheets_objs.insert(0,w.worksheets_objs.pop(w.worksheets_objs.index(summary)))
 summary.activate();summary.set_first_sheet()
 book.close()
 # Machine-readable companion and a lightweight audit of the delivered workbook.
 with (ROOT/'summary.csv').open('w') as f:
  writer=csv.writer(f);writer.writerow(['model',*q.keys()])
  for m in ('qwen','glm'):writer.writerow([m,*[json.dumps(v) if isinstance(v,(dict,list)) else v for v in a['models'][m].values()]])
 import openpyxl
 wb=openpyxl.load_workbook(ROOT/'AgentBenchmark_Qwen38_vs_GLM53.xlsx',data_only=True)
 assert wb.sheetnames[0]=='Summary';assert wb['Runs'].max_row==148;assert len(wb.sheetnames)>=12
 assert wb['Summary']['B8'].value is not None and wb['Summary']['C8'].value is not None
 assert len(wb['Summary']._charts)==2
 with zipfile.ZipFile(ROOT/'AgentBenchmark_Qwen38_vs_GLM53.xlsx') as z:assert z.testzip() is None
 dump={'worksheets':wb.sheetnames,'run_rows':wb['Runs'].max_row-4,'checks_rows':wb['Checks'].max_row-4,'tool_call_rows':wb['ToolCalls'].max_row-4,'formula_cached_values_present':True,'xlsx_zip_integrity':'OK'}
 (ROOT/'workbook-validation.json').write_text(json.dumps(dump,indent=2));print(json.dumps(a['models'],indent=2));print('WORKBOOK_OK',ROOT/'AgentBenchmark_Qwen38_vs_GLM53.xlsx')

if __name__=='__main__':create_report()
